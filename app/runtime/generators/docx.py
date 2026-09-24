import hashlib
import json
from datetime import UTC, datetime
from pathlib import Path

from docx import Document
from pydantic import BaseModel, ConfigDict, Field, ValidationError

from app.errors import AppError

MAX_SPEC_BYTES = 1024 * 1024


class _Strict(BaseModel):
    model_config = ConfigDict(extra="forbid")


class TableSpec(_Strict):
    columns: list[str] = Field(max_length=100)
    rows: list[list[str]] = Field(max_length=1000)


class SectionSpec(_Strict):
    heading: str = Field(max_length=500)
    paragraphs: list[str] = Field(default_factory=list, max_length=1000)
    table: TableSpec | None = None
    citations: list[str] = Field(default_factory=list, max_length=1000)


class EvidenceSpec(_Strict):
    chunk_id: str
    document_id: str
    page: int = Field(ge=1)
    section: str
    source_hash: str


class CalculationSpec(_Strict):
    name: str
    formula: str
    inputs: dict[str, str]
    result: str


class DocxSpec(_Strict):
    title: str = Field(max_length=1000)
    metadata: dict[str, str] = Field(default_factory=dict, max_length=100)
    sections: list[SectionSpec] = Field(default_factory=list, max_length=200)
    evidence: list[EvidenceSpec] = Field(default_factory=list, max_length=1000)
    calculations: list[CalculationSpec] = Field(default_factory=list, max_length=1000)
    assumptions: list[str] = Field(default_factory=list, max_length=1000)
    recommendation: str = Field(max_length=10000)


def _spec(value: dict) -> DocxSpec:
    try:
        encoded = json.dumps(value, sort_keys=True).encode()
        if len(encoded) > MAX_SPEC_BYTES:
            raise ValueError("spec exceeds 1 MiB")
        return DocxSpec.model_validate(value)
    except (ValidationError, ValueError, TypeError) as exc:
        raise AppError("VALIDATION_ERROR", 422, "Invalid DOCX specification") from exc


def generate(path: Path, value: dict) -> str:
    spec = _spec(value)
    canonical = json.dumps(spec.model_dump(), sort_keys=True, separators=(",", ":"))
    document = Document()
    document.add_heading(spec.title, 0)
    metadata = document.add_table(rows=1, cols=2)
    metadata.rows[0].cells[0].text, metadata.rows[0].cells[1].text = "Key", "Value"
    for key, value in spec.metadata.items():
        cells = metadata.add_row().cells
        cells[0].text, cells[1].text = key, value
    for section in spec.sections:
        document.add_heading(section.heading, level=1)
        for paragraph in section.paragraphs:
            document.add_paragraph(paragraph)
        if section.table:
            table = document.add_table(rows=1, cols=len(section.table.columns))
            for cell, value in zip(table.rows[0].cells, section.table.columns, strict=True):
                cell.text = value
            for row in section.table.rows:
                if len(row) != len(section.table.columns):
                    raise AppError(
                        "VALIDATION_ERROR", 422, "Table row length does not match columns"
                    )
                for cell, value in zip(table.add_row().cells, row, strict=True):
                    cell.text = value
        if section.citations:
            document.add_paragraph(" ".join(f"[{item}]" for item in section.citations))
    document.add_heading("Evidence", level=1)
    evidence = document.add_table(rows=1, cols=5)
    for cell, name in zip(
        evidence.rows[0].cells,
        ["chunk_id", "document_id", "page", "section", "source_hash"],
        strict=True,
    ):
        cell.text = name
    for item in spec.evidence:
        for cell, value in zip(
            evidence.add_row().cells,
            [item.chunk_id, item.document_id, str(item.page), item.section, item.source_hash],
            strict=True,
        ):
            cell.text = value
    document.add_heading("Calculations", level=1)
    for item in spec.calculations:
        document.add_paragraph(
            f"{item.name}: {item.formula}; inputs={item.inputs}; result={item.result}"
        )
    document.add_heading("Assumptions", level=1)
    for item in spec.assumptions:
        document.add_paragraph(item, style="List Bullet")
    document.add_heading("Recommendation", level=1)
    document.add_paragraph(spec.recommendation)
    document.add_heading("Approval", level=1)
    document.add_paragraph(
        "Approved by: ____________________\nDate: ____________________\nSignature: ____________________"
    )
    footer = document.sections[0].footer.paragraphs[0]
    footer.text = f"Generated {datetime.now(UTC).isoformat().replace('+00:00', 'Z')} | content_hash {hashlib.sha256(canonical.encode()).hexdigest()}"
    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(path)
    return hashlib.sha256(canonical.encode()).hexdigest()
