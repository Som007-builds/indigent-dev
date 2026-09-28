import hashlib
import json
import uuid
import zipfile
from datetime import UTC, datetime
from pathlib import Path
from typing import Any

from app.contracts.models import ArtifactManifest, TaskContext
from app.core.audit import AuditLoggerImpl
from app.core.repo import Repository
from app.core.workspace import safe_join
from app.errors import AppError, ArtifactCorruptError, PathRejectedError


def _now() -> str:
    return datetime.now(UTC).isoformat().replace("+00:00", "Z")


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as source:
        while chunk := source.read(1024 * 1024):
            digest.update(chunk)
    return digest.hexdigest()


def annotate_size(record: dict[str, Any]) -> dict[str, Any]:
    """Attach the artifact's real on-disk size to a manifest or ``artifacts`` row.

    No artifact size is ever persisted: the ``artifacts`` table (G10) has no size
    column and ``ArtifactManifest`` (G8) has no size field. The file is on disk though,
    so the size is knowable at read time -- and because it was not, the frontend had
    nothing to display and rendered the literal string "Validated deliverable" in a
    size field (Frontend-fix.md item 3.2).

    ``None`` when the file is missing or unreadable. Absent means *not reported*, which
    is deliberately not the same as zero: reporting 0 bytes for a file that has been
    deleted would be a different kind of lie.

    The path comes from our own database rather than from a request, and ``register``
    already constrained it under the task's ``outputs/``, so stat-ing it is safe. This
    reads metadata only -- it never opens the file, which is what
    ``verify_integrity`` and the download route are for.
    """
    enriched = dict(record)
    path = Path(str(record.get("path") or ""))
    try:
        enriched["size_bytes"] = path.stat().st_size if path.is_file() else None
    except OSError:
        enriched["size_bytes"] = None
    return enriched


class ArtifactStoreImpl:
    def __init__(self, repo: Repository, audit: AuditLoggerImpl) -> None:
        self.repo, self.audit = repo, audit

    async def register(
        self,
        ctx: TaskContext,
        artifact_type: str,
        path: str,
        source_evidence_ids: list[str] | None = None,
        metadata: dict | None = None,
    ) -> ArtifactManifest:
        outputs = safe_join(ctx.workspace, "outputs")
        candidate = Path(path)
        try:
            relative = candidate.resolve().relative_to(outputs.resolve())
        except ValueError as exc:
            raise PathRejectedError("Artifact must be within task outputs") from exc
        artifact_path = safe_join(outputs, relative)
        if not artifact_path.is_file():
            raise AppError("NOT_FOUND", 404, "Artifact file not found")
        manifest = ArtifactManifest(
            artifact_id=str(uuid.uuid4()),
            task_id=ctx.task_id,
            artifact_type=artifact_type,
            path=str(artifact_path),
            created_at=_now(),
            artifact_hash=_sha256(artifact_path),
            source_evidence_ids=source_evidence_ids or [],
            metadata=metadata or {},
        )
        manifest_path = safe_join(outputs, f"{manifest.artifact_id}.manifest.json")
        manifest_path.write_text(manifest.model_dump_json(indent=2), encoding="utf-8")
        await self.repo.insert_artifact(manifest)
        await self.audit.emit(
            "ARTIFACT_CREATED",
            "artifacts",
            "register",
            "ok",
            ctx.task_id,
            {"artifact_id": manifest.artifact_id, "artifact_type": artifact_type},
        )
        return manifest

    async def set_verification_status(self, artifact_id: str, status: str) -> None:
        if status not in {"pending", "passed", "failed"}:
            raise ValueError("invalid verification status")
        manifest = await self.get(artifact_id)
        async with self.repo.db.connection() as conn:
            await conn.execute(
                "UPDATE artifacts SET verification_status=? WHERE artifact_id=?",
                (status, artifact_id),
            )
            await conn.commit()
        manifest.verification_status = status
        manifest_path = safe_join(Path(manifest.path).parent, f"{artifact_id}.manifest.json")
        manifest_path.write_text(manifest.model_dump_json(indent=2), encoding="utf-8")

    async def get(self, artifact_id: str) -> ArtifactManifest:
        row = await self.repo.get_artifact(artifact_id)
        if row is None:
            raise AppError("NOT_FOUND", 404, "Artifact not found")
        return ArtifactManifest(
            **{
                **row,
                "source_evidence_ids": json.loads(row["source_evidence_ids"]),
                "metadata": json.loads(row["metadata"]),
            }
        )

    async def verify_integrity(self, artifact_id: str) -> bool:
        manifest = await self.get(artifact_id)
        path = Path(manifest.path)
        valid = path.is_file() and _sha256(path) == manifest.artifact_hash
        if not valid:
            await self.audit.emit(
                "ARTIFACT_CORRUPT",
                "artifacts",
                "integrity check",
                "error",
                manifest.task_id,
                {"artifact_id": artifact_id},
            )
        return valid

    async def package_code(self, ctx: TaskContext) -> ArtifactManifest:
        outputs, code = safe_join(ctx.workspace, "outputs"), safe_join(ctx.workspace, "code")
        archive = safe_join(outputs, "code-package.zip")
        with zipfile.ZipFile(archive, "w", zipfile.ZIP_DEFLATED) as bundle:
            for root, prefix in ((code, "code"), (outputs, "outputs")):
                patterns = ("**/*",) if prefix == "code" else ("test_results*.json",)
                for pattern in patterns:
                    for item in root.glob(pattern):
                        if item.is_file() and item != archive:
                            bundle.write(item, Path(prefix) / item.relative_to(root))
        return await self.register(ctx, "code_package", str(archive))

    async def create_stub_docx(self, ctx: TaskContext) -> ArtifactManifest:
        from docx import Document

        outputs = safe_join(ctx.workspace, "outputs")
        outputs.mkdir(parents=True, exist_ok=True)
        path = outputs / "stub-report.docx"
        document = Document()
        document.add_heading("Indigent Verification Report", level=1)
        document.add_paragraph("Industrial AI-Harness Sovereign Report")
        document.add_paragraph(f"Task ID: {ctx.task_id}")
        document.save(path)
        return await self.register(
            ctx,
            "docx",
            str(path),
            metadata={"title": "Verification Report", "format": "docx"},
        )

    async def require_intact(self, artifact_id: str) -> ArtifactManifest:
        manifest = await self.get(artifact_id)
        if not await self.verify_integrity(artifact_id):
            raise ArtifactCorruptError()
        return manifest

