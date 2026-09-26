from __future__ import annotations

import ast
import hashlib
import json
import zipfile
from dataclasses import dataclass
from pathlib import Path
from typing import TYPE_CHECKING, Any, Literal

from pydantic import ValidationError

from app.contracts.models import ArtifactManifest

if TYPE_CHECKING:
    from app.rag.models import ClaimProvenance, EvidenceItem, VerificationProvenance
    from app.rag.verification.citation import CitationVerificationResult


ValidationStatus = Literal["VALID", "INVALID", "UNVERIFIED"]


@dataclass(frozen=True)
class ArtifactValidationContext:
    claims: tuple[ClaimProvenance, ...] = ()
    evidence: tuple[EvidenceItem, ...] = ()
    citation_results: tuple[CitationVerificationResult, ...] = ()
    verifications: tuple[VerificationProvenance, ...] = ()
    expected_fields: tuple[str, ...] = ()
    calculations: tuple[dict[str, Any], ...] = ()
    provenance: tuple[dict[str, Any], ...] = ()


@dataclass(frozen=True)
class ArtifactValidationResult:
    status: ValidationStatus
    valid: bool
    artifact_hash: str | None
    checks: tuple[dict[str, Any], ...] = ()
    provenance: tuple[dict[str, Any], ...] = ()


class SemanticArtifactValidator:
    def __init__(self, citation_verifier: Any | None = None) -> None:
        self.citation_verifier = citation_verifier

    async def validate(
        self,
        artifact: ArtifactManifest,
        context: ArtifactValidationContext | None = None,
    ) -> ArtifactValidationResult:
        context = context or ArtifactValidationContext()
        checks: list[dict[str, Any]] = []
        try:
            path = Path(artifact.path)
            if not path.is_file():
                return self._result("INVALID", None, [{"name": "exists", "passed": False, "reason": "ARTIFACT_MISSING"}], context)
            try:
                content = path.read_bytes()
            except OSError:
                return self._result("INVALID", None, [{"name": "readable", "passed": False, "reason": "ARTIFACT_UNREADABLE"}], context)
            actual_hash = hashlib.sha256(content).hexdigest()
            checks.append({"name": "exists_readable", "passed": True})
            if artifact.metadata is None or not isinstance(artifact.metadata, dict):
                return self._result("INVALID", actual_hash, [{"name": "manifest_metadata", "passed": False, "reason": "MISSING_MANIFEST_METADATA"}], context)
            if not artifact.artifact_id or not artifact.task_id or not artifact.created_at:
                return self._result("INVALID", actual_hash, [{"name": "manifest_required_fields", "passed": False, "reason": "MISSING_MANIFEST_FIELDS"}], context)
            checks.append({"name": "manifest_required_fields", "passed": True})
            structural, unverified = self._structural(artifact, path, content, context)
            checks.extend(structural)
            if any(check.get("passed") is False for check in structural):
                return self._result("INVALID", actual_hash, checks, context)
            semantic, semantic_unverified = self._semantic(artifact, context)
            checks.extend(semantic)
            status: ValidationStatus = "INVALID" if any(item.get("passed") is False for item in semantic) else (
                "UNVERIFIED" if unverified or semantic_unverified else "VALID"
            )
            return self._result(status, actual_hash, checks, context)
        except NameError:
            raise
        except Exception as error:
            checks.append({"name": "validator_exception", "passed": False, "reason": "VALIDATOR_ERROR", "detail": type(error).__name__})
            return self._result("INVALID", None, checks, context)

    def _structural(
        self, artifact: ArtifactManifest, path: Path, content: bytes, context: ArtifactValidationContext
    ) -> tuple[list[dict[str, Any]], bool]:
        checks: list[dict[str, Any]] = []
        expected_ext = {
            "docx": ".docx",
            "xlsx": ".xlsx",
            "pptx": ".pptx",
            "graph_json": ".json",
            "code_package": ".zip",
            "pid_overlay": path.suffix.lower(),
        }.get(artifact.artifact_type)
        if expected_ext is None or path.suffix.lower() != expected_ext:
            return [{"name": "artifact_type_extension", "passed": False, "reason": "ARTIFACT_TYPE_MISMATCH"}], False
        checks.append({"name": "artifact_type_extension", "passed": True})
        unverified = False
        if artifact.artifact_type == "docx":
            try:
                from docx import Document as DocxDocument
                doc = DocxDocument(path)
                if not doc.paragraphs and not doc.tables:
                    raise ValueError("empty document")
                checks.append({"name": "docx_readable", "passed": True})
            except ImportError:
                checks.append({"name": "docx_readable", "passed": None, "reason": "DOCX_VALIDATOR_UNAVAILABLE"})
                unverified = True
            except Exception:
                checks.append({"name": "docx_readable", "passed": False, "reason": "MALFORMED_DOCX"})
        elif artifact.artifact_type == "xlsx":
            try:
                from openpyxl import load_workbook
                workbook = load_workbook(path, read_only=True, data_only=False)
                valid = bool(workbook.sheetnames)
                workbook.close()
                checks.append({"name": "xlsx_readable", "passed": valid, "reason": None if valid else "EMPTY_XLSX"})
            except ImportError:
                checks.append({"name": "xlsx_readable", "passed": None, "reason": "XLSX_VALIDATOR_UNAVAILABLE"})
                unverified = True
            except Exception:
                checks.append({"name": "xlsx_readable", "passed": False, "reason": "MALFORMED_XLSX"})
        elif artifact.artifact_type in {"pptx", "code_package", "pid_overlay"}:
            if artifact.artifact_type == "pptx":
                checks.append({"name": "pptx_validation", "passed": None, "reason": "UNSUPPORTED_VALIDATOR"})
                unverified = True
            else:
                try:
                    with zipfile.ZipFile(path) as archive:
                        corrupt_member = archive.testzip()
                    valid = corrupt_member is None
                    checks.append({"name": "zip_integrity", "passed": valid, "reason": None if valid else "CORRUPT_ARCHIVE"})
                except (zipfile.BadZipFile, OSError):
                    checks.append({"name": "zip_integrity", "passed": False, "reason": "MALFORMED_ARCHIVE"})
                if artifact.artifact_type == "pid_overlay":
                    unverified = True
        elif artifact.artifact_type == "graph_json":
            try:
                value = json.loads(content)
                if not isinstance(value, dict) or not {"nodes", "edges"}.issubset(value):
                    raise ValueError("graph fields missing")
                if not isinstance(value["nodes"], list) or not isinstance(value["edges"], list):
                    raise ValueError("graph arrays invalid")
                checks.append({"name": "graph_json_schema", "passed": True})
            except (json.JSONDecodeError, ValueError, TypeError):
                checks.append({"name": "graph_json_schema", "passed": False, "reason": "MALFORMED_GRAPH_JSON"})
        if artifact.artifact_type == "docx":
            try:
                from app.runtime.generators.docx import DocxSpec

                spec_value = artifact.metadata.get("spec")
                if spec_value is not None:
                    spec = DocxSpec.model_validate(spec_value)
                    checks.extend(self._validate_spec_fields(spec.model_dump(), artifact.metadata, context))
                    checks.extend(self._validate_calculations(spec.calculations, context.calculations))
                elif context.expected_fields:
                    unverified = True
                    checks.append({"name": "required_fields", "passed": None, "reason": "SPEC_UNAVAILABLE"})
            except ValidationError:
                checks.append({"name": "docx_spec", "passed": False, "reason": "INVALID_DOCX_SPEC"})
            except ImportError:
                unverified = True
                checks.append({"name": "docx_spec", "passed": None, "reason": "DOCX_VALIDATOR_UNAVAILABLE"})
        elif artifact.artifact_type == "xlsx":
            spec_value = artifact.metadata.get("spec")
            if spec_value is not None:
                try:
                    from app.runtime.generators.xlsx import XlsxSpec

                    spec = XlsxSpec.model_validate(spec_value)
                    checks.extend(self._validate_calculations([], context.calculations))
                    formulas = [formula for sheet in spec.sheets for formula in sheet.formulas]
                    if formulas:
                        unverified = True
                        checks.append({"name": "spreadsheet_formulas", "passed": None, "reason": "FORMULAS_NOT_RECOMPUTED"})
                except ValidationError:
                    checks.append({"name": "xlsx_spec", "passed": False, "reason": "INVALID_XLSX_SPEC"})
                except ImportError:
                    unverified = True
                    checks.append({"name": "xlsx_spec", "passed": None, "reason": "XLSX_VALIDATOR_UNAVAILABLE"})
        return checks, unverified

    def _validate_spec_fields(
        self, spec: dict[str, Any], metadata: dict[str, Any], context: ArtifactValidationContext
    ) -> list[dict[str, Any]]:
        available = set(spec) | set(metadata)
        missing = sorted(field for field in context.expected_fields if field not in available)
        return [{"name": "required_fields", "passed": not missing, "reason": "MISSING_REQUIRED_FIELDS" if missing else None, "missing": missing}]

    def _semantic(
        self, artifact: ArtifactManifest, context: ArtifactValidationContext
    ) -> tuple[list[dict[str, Any]], bool]:
        checks: list[dict[str, Any]] = []
        unverified = False
        evidence = {item.chunk_id: item for item in context.evidence}
        verification_by_id = {item.verification_id: item for item in context.verifications}
        citation_by_chunk: dict[str, list[CitationVerificationResult]] = {}
        for result in context.citation_results:
            citation_by_chunk.setdefault(result.evidence.chunk_id, []).append(result)
        for claim in context.claims:
            if not claim.evidence_chunk_ids:
                checks.append({"name": "claim_evidence", "passed": False, "reason": "CLAIM_MISSING_EVIDENCE", "claim_id": claim.claim_id})
                continue
            for chunk_id in claim.evidence_chunk_ids:
                item = evidence.get(chunk_id)
                if item is None:
                    checks.append({"name": "claim_evidence", "passed": False, "reason": "EVIDENCE_NOT_FOUND", "claim_id": claim.claim_id, "chunk_id": chunk_id})
                    continue
                citation = next((r for r in citation_by_chunk.get(chunk_id, []) if r.claim_id == claim.claim_id), None)
                verification = verification_by_id.get(claim.verification_id or "")
                if citation is not None:
                    passed = citation.verified
                    checks.append({"name": "citation_verification", "passed": passed, "reason": None if passed else "CITATION_FAILED", "claim_id": claim.claim_id, "chunk_id": chunk_id})
                elif verification is not None:
                    checks.append({"name": "citation_verification", "passed": verification.passed, "reason": verification.detail, "claim_id": claim.claim_id, "chunk_id": chunk_id})
                else:
                    unverified = True
                    checks.append({"name": "citation_verification", "passed": None, "reason": "CITATION_UNVERIFIED", "claim_id": claim.claim_id, "chunk_id": chunk_id})
        checks.extend(self._validate_calculations([], context.calculations))
        if any(check.get("passed") is None for check in checks):
            unverified = True
        for item in context.provenance:
            if item.get("artifact_id") not in {None, artifact.artifact_id}:
                checks.append({"name": "provenance_artifact", "passed": False, "reason": "ARTIFACT_REFERENCE_MISMATCH"})
        return checks, unverified

    def _validate_calculations(
        self, artifact_calculations: list[Any], context_calculations: tuple[dict[str, Any], ...]
    ) -> list[dict[str, Any]]:
        checks: list[dict[str, Any]] = []
        for index, calculation in enumerate([*artifact_calculations, *context_calculations]):
            if not isinstance(calculation, dict):
                calculation = calculation.model_dump()
            name = calculation.get("name", f"calculation-{index}")
            formula = calculation.get("formula")
            inputs = calculation.get("inputs")
            expected = calculation.get("result")
            if not isinstance(formula, str) or not isinstance(inputs, dict) or expected is None:
                checks.append({"name": "calculation", "passed": None, "reason": "CALCULATION_UNVERIFIED", "calculation": name})
                continue
            computed = _safe_arithmetic(formula, inputs)
            if computed is None:
                checks.append({"name": "calculation", "passed": None, "reason": "CALCULATION_UNVERIFIED", "calculation": name})
                continue
            try:
                matches = abs(computed - float(expected)) <= max(1e-9, abs(computed) * 1e-9)
            except (TypeError, ValueError):
                matches = False
            checks.append({"name": "calculation", "passed": matches, "reason": None if matches else "CALCULATION_MISMATCH", "calculation": name, "computed": computed})
        return checks

    def _result(
        self,
        status: ValidationStatus,
        artifact_hash: str | None,
        checks: list[dict[str, Any]],
        context: ArtifactValidationContext,
    ) -> ArtifactValidationResult:
        provenance = list(context.provenance)
        provenance.extend(
            {
                "document_id": item.document_id,
                "chunk_id": item.chunk_id,
                "source_hash": item.source_hash,
                "claim_id": claim.claim_id,
                "verification_id": claim.verification_id,
                "artifact_id": claim.artifact_id,
            }
            for claim in context.claims
            for item in context.evidence
            if item.chunk_id in claim.evidence_chunk_ids
        )
        for item in provenance:
            item.setdefault("artifact_id", None)
            item.setdefault("artifact_hash", artifact_hash)
        return ArtifactValidationResult(status, status == "VALID", artifact_hash, tuple(checks), tuple(provenance))


def _safe_arithmetic(formula: str, inputs: dict[str, Any]) -> float | None:
    try:
        values = {name: float(value) for name, value in inputs.items()}
        expression = formula
        for name, value in values.items():
            expression = expression.replace(name, repr(value))
        tree = ast.parse(expression, mode="eval")
        allowed = (ast.Expression, ast.BinOp, ast.UnaryOp, ast.Constant, ast.Add, ast.Sub, ast.Mult, ast.Div, ast.Pow, ast.USub, ast.UAdd, ast.Load)
        if any(not isinstance(node, allowed) for node in ast.walk(tree)):
            return None
        result = eval(compile(tree, "<calculation>", "eval"), {"__builtins__": {}}, {})
        return float(result) if isinstance(result, (int, float)) else None
    except (ArithmeticError, SyntaxError, TypeError, ValueError, NameError):
        return None
