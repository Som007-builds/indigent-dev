import json
import zipfile

import pytest
from docx import Document
from httpx import ASGITransport, AsyncClient
from openpyxl import load_workbook

from app.config import Settings
from app.contracts.models import TaskContext
from app.core.artifacts import ArtifactStoreImpl
from app.core.audit import AuditLoggerImpl
from app.core.db import Database
from app.core.repo import Repository
from app.core.workspace import create_workspace
from app.errors import AppError, PathRejectedError
from app.main import create_app
from app.runtime.generators.docx import generate as generate_docx
from app.runtime.generators.xlsx import generate as generate_xlsx


async def artifact_store(tmp_path):
    settings = Settings(data_dir=tmp_path)
    repo = Repository(Database(tmp_path / "db.sqlite"))
    await repo.db.initialize()
    audit = AuditLoggerImpl(repo, settings)
    await repo.create_task("task", "request", "local")
    workspace = create_workspace(tmp_path / "workspaces", "task")
    return (
        ArtifactStoreImpl(repo, audit),
        repo,
        TaskContext(
            task_id="task", task_type="inspection", workspace=workspace, inference_mode="local"
        ),
    )


@pytest.mark.asyncio
async def test_register_manifest_and_package_code(tmp_path):
    store, repo, ctx = await artifact_store(tmp_path)
    output = ctx.workspace / "outputs" / "report.txt"
    output.write_text("artifact", encoding="utf-8")
    manifest = await store.register(ctx, "docx", str(output))
    manifest_file = ctx.workspace / "outputs" / f"{manifest.artifact_id}.manifest.json"
    assert manifest_file.exists()
    assert json.loads((await repo.get_artifact(manifest.artifact_id))["metadata"]) == {}
    assert json.loads(manifest_file.read_text(encoding="utf-8"))["artifact_hash"] == manifest.artifact_hash
    assert await store.verify_integrity(manifest.artifact_id)
    await store.set_verification_status(manifest.artifact_id, "passed")
    assert json.loads(manifest_file.read_text(encoding="utf-8"))["verification_status"] == "passed"
    (ctx.workspace / "code" / "main.py").write_text("print('ok')", encoding="utf-8")
    (ctx.workspace / "outputs" / "test_results.json").write_text("{}", encoding="utf-8")
    package = await store.package_code(ctx)
    with zipfile.ZipFile(package.path) as archive:
        assert {"code/main.py", "outputs/test_results.json"} <= set(archive.namelist())


@pytest.mark.asyncio
async def test_artifact_path_and_tamper_download_are_rejected(tmp_path):
    store, repo, ctx = await artifact_store(tmp_path)
    outside = tmp_path / "outside.txt"
    outside.write_text("x", encoding="utf-8")
    with pytest.raises(PathRejectedError):
        await store.register(ctx, "docx", str(outside))
    output = ctx.workspace / "outputs" / "report.docx"
    output.write_text("safe", encoding="utf-8")
    manifest = await store.register(ctx, "docx", str(output))
    output.write_text("tampered", encoding="utf-8")
    app = create_app(Settings(data_dir=tmp_path))
    async with app.router.lifespan_context(app):
        async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as client:
            response = await client.get(f"/api/artifacts/{manifest.artifact_id}?download=1")
    assert response.status_code == 500
    assert "ARTIFACT_CORRUPT" in [item["category"] for item in await repo.list_audit("task")]


@pytest.mark.asyncio
async def test_missing_artifact_download_is_not_reported_as_corruption(tmp_path):
    store, _, ctx = await artifact_store(tmp_path)
    output = ctx.workspace / "outputs" / "report.docx"
    output.write_text("safe", encoding="utf-8")
    manifest = await store.register(ctx, "docx", str(output))
    output.unlink()
    app = create_app(Settings(data_dir=tmp_path))
    async with app.router.lifespan_context(app):
        async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as client:
            response = await client.get(f"/api/artifacts/{manifest.artifact_id}?download=1")
    assert response.status_code == 404


def test_docx_and_xlsx_generators_reopen_and_reject_bad_specs(tmp_path):
    docx_path = tmp_path / "report.docx"
    generate_docx(
        docx_path,
        {
            "title": "Report",
            "metadata": {"owner": "PSU"},
            "sections": [{"heading": "Findings", "paragraphs": ["Safe"], "citations": ["chunk-1"]}],
            "evidence": [
                {
                    "chunk_id": "chunk-1",
                    "document_id": "doc",
                    "page": 1,
                    "section": "A",
                    "source_hash": "hash",
                }
            ],
            "calculations": [
                {"name": "total", "formula": "1+1", "inputs": {"a": "1"}, "result": "2"}
            ],
            "assumptions": ["stable"],
            "recommendation": "Proceed",
        },
    )
    text = "\n".join(paragraph.text for paragraph in Document(docx_path).paragraphs)
    assert all(
        item in text
        for item in [
            "Report",
            "Findings",
            "[chunk-1]",
            "Calculations",
            "Assumptions",
            "Recommendation",
            "Approval",
        ]
    )
    xlsx_path = tmp_path / "report.xlsx"
    generate_xlsx(
        xlsx_path,
        {
            "sheets": [
                {
                    "name": "Data",
                    "columns": ["n"],
                    "rows": [[1]],
                    "formulas": [{"cell": "B2", "formula": "=SUM(A2:A2)"}],
                }
            ]
        },
    )
    assert load_workbook(xlsx_path, data_only=False)["Data"]["B2"].value == "=SUM(A2:A2)"
    with pytest.raises(AppError):
        generate_xlsx(tmp_path / "bad.xlsx", {"sheets": [{"name": "bad/name", "columns": []}]})
    assert not (tmp_path / "bad.xlsx").exists()
    with pytest.raises(AppError):
        generate_xlsx(
            tmp_path / "bad-cell.xlsx",
            {"sheets": [{"name": "Data", "columns": [], "formulas": [{"cell": "!", "formula": "=1"}]}]},
        )
    assert not (tmp_path / "bad-cell.xlsx").exists()
